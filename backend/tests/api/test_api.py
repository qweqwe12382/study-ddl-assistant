from io import BytesIO
from datetime import date, timedelta

import fitz
from docx import Document
from PIL import Image, ImageDraw
from icalendar import Calendar


class _BrokenProvider:
    name = "broken-test-provider"

    def extract(self, _text, _filename):
        return "{not-json"


class _TimeoutProvider:
    name = "timeout-test-provider"

    def extract(self, _text, _filename):
        from app.services.llm_provider import ProviderError

        raise ProviderError("LLM_TIMEOUT", "模型调用超时，请稍后重试")


class _NetworkFailureProvider:
    name = "network-failure-test-provider"

    def extract(self, _text, _filename):
        from app.services.llm_provider import ProviderError

        raise ProviderError("LLM_REQUEST_FAILED", "模型服务连接失败，请稍后重试")


def _upload(client, filename, content, content_type, **data):
    return client.post(
        "/api/materials/upload",
        data=data,
        files={"files": (filename, content, content_type)},
    )


def test_health_check(client):
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"
    assert response.json()["database"] == "ok"


def test_m7_health_check_returns_request_id(client):
    response = client.get("/api/health", headers={"X-Request-ID": "m7-health-check"})

    assert response.status_code == 200
    assert response.headers["X-Request-ID"] == "m7-health-check"


def test_request_id_rejects_unbounded_client_values(client):
    response = client.get("/api/health", headers={"X-Request-ID": "x" * 200})

    assert response.status_code == 200
    assert response.headers["X-Request-ID"] != "x" * 200
    assert len(response.headers["X-Request-ID"]) == 16


def test_m7_health_check_reports_database_failure(client):
    from app.database import get_system_db
    from app.main import app
    from sqlalchemy.exc import SQLAlchemyError

    class BrokenSession:
        def execute(self, _statement):
            raise SQLAlchemyError("database unavailable")

    def broken_db():
        yield BrokenSession()

    app.dependency_overrides[get_system_db] = broken_db
    try:
        response = client.get("/api/health")
    finally:
        app.dependency_overrides.pop(get_system_db, None)

    assert response.status_code == 503
    assert response.json()["error"]["code"] == "DATABASE_UNAVAILABLE"


def test_llm_settings_endpoint_removed(client):
    assert client.get("/api/settings/llm").status_code == 404
    assert client.put("/api/settings/llm", json={"model": "unused"}).status_code == 404


def test_m7_reset_clears_business_data_and_uploaded_files(client, isolated_upload_dir):
    course = client.post("/api/courses", json={"name": "测试课程"}).json()
    material = client.post(
        "/api/materials",
        json={
            "course_id": course["id"],
            "original_filename": "测试通知.txt",
            "extracted_text": "完成测试任务。",
        },
    ).json()
    client.post(
        "/api/tasks",
        json={
            "course_id": course["id"],
            "material_id": material["id"],
            "name": "完成测试任务",
        },
    )
    plan = client.post(
        "/api/study-plans/generate",
        json={
            "course_id": course["id"],
            "exam_date": (date.today() + timedelta(days=3)).isoformat(),
            "daily_minutes": 60,
        },
    )
    assert plan.status_code == 201

    stale_file = isolated_upload_dir / "stale-upload.txt"
    stale_file.write_text("stale", encoding="utf-8")

    response = client.post("/api/dev/reset")

    assert response.status_code == 200
    assert response.json()["deleted"] == {
        "study_plans": 1,
        "tasks": 1,
        "materials": 1,
        "courses": 1,
    }
    assert response.json()["deleted_files"] == 1
    assert not stale_file.exists()
    assert client.get("/api/courses").json() == []
    assert client.get("/api/materials").json() == []
    assert client.get("/api/tasks").json() == []
    assert client.get("/api/study-plans").json() == []


def test_m7_reset_is_disabled_outside_development(client):
    from app.config import settings

    previous = settings.app_env
    object.__setattr__(settings, "app_env", "production")
    try:
        response = client.post("/api/dev/reset")
    finally:
        object.__setattr__(settings, "app_env", previous)

    assert response.status_code == 404
    assert response.json()["error"]["code"] == "RESET_DISABLED"


def test_database_tables_are_initialized(client):
    response = client.get("/api/courses")
    assert response.status_code == 200
    assert response.json() == []


def test_course_crud(client):
    created = client.post("/api/courses", json={"name": "数据结构", "semester": "2026 秋"})
    assert created.status_code == 201
    course = created.json()
    assert course["name"] == "数据结构"

    updated = client.patch(f"/api/courses/{course['id']}", json={"teacher": "王老师"})
    assert updated.status_code == 200
    assert updated.json()["teacher"] == "王老师"

    listed = client.get("/api/courses")
    assert listed.status_code == 200
    assert len(listed.json()) == 1


def test_text_fields_are_trimmed_and_whitespace_names_are_rejected(client):
    created = client.post("/api/courses", json={"name": "  数据结构  ", "teacher": "  王老师  "})
    invalid_course = client.post("/api/courses", json={"name": "   "})
    invalid_task = client.post("/api/tasks", json={"name": "\t "})

    assert created.status_code == 201
    assert created.json()["name"] == "数据结构"
    assert created.json()["teacher"] == "王老师"
    assert invalid_course.status_code == 422
    assert invalid_task.status_code == 422


def test_material_and_task_crud(client):
    course = client.post("/api/courses", json={"name": "数据库原理"}).json()
    material = client.post(
        "/api/materials",
        json={
            "course_id": course["id"],
            "original_filename": "课程通知.txt",
            "file_type": "txt",
            "tags": ["作业", "DDL"],
        },
    )
    assert material.status_code == 201
    material_id = material.json()["id"]

    task = client.post(
        "/api/tasks",
        json={
            "course_id": course["id"],
            "material_id": material_id,
            "name": "完成实验一",
            "priority": 4,
        },
    )
    assert task.status_code == 201
    task_id = task.json()["id"]

    completed = client.post(f"/api/tasks/{task_id}/complete")
    assert completed.status_code == 200
    assert completed.json()["status"] == "completed"


def test_m5_material_search_filters_and_returns_match_evidence(client):
    course = client.post("/api/courses", json={"name": "数据库原理"}).json()
    client.post(
        "/api/materials",
        json={
            "course_id": course["id"],
            "original_filename": "实验通知.md",
            "material_type": "作业要求",
            "tags": ["DDL", "重点"],
            "summary": "本周实验提交说明",
            "extracted_text": "请完成 SQL 查询实验，截止时间为周五。",
        },
    )
    client.post(
        "/api/materials",
        json={
            "original_filename": "阅读材料.txt",
            "material_type": "教材或阅读材料",
            "extracted_text": "关系模型基础阅读。",
        },
    )

    response = client.get(
        "/api/materials",
        params={"q": "SQL", "course_id": course["id"], "material_type": "作业要求", "tag": "ddl"},
    )

    assert response.status_code == 200
    rows = response.json()
    assert len(rows) == 1
    assert rows[0]["original_filename"] == "实验通知.md"
    assert rows[0]["matched_fields"] == ["content"]
    assert "SQL" in rows[0]["match_snippets"][0]


def test_m5_dashboard_returns_task_windows_and_recent_materials(client):
    course = client.post("/api/courses", json={"name": "操作系统"}).json()
    material = client.post(
        "/api/materials",
        json={"course_id": course["id"], "original_filename": "课程通知.txt", "extracted_text": "DDL"},
    ).json()
    client.post(
        "/api/tasks",
        json={
            "course_id": course["id"],
            "material_id": material["id"],
            "name": "未来任务",
            "due_at": "2099-01-02T12:00:00",
            "priority": 5,
        },
    )
    client.post(
        "/api/tasks",
        json={"name": "逾期任务", "due_at": "2020-01-02T12:00:00", "status": "not_started"},
    )

    response = client.get("/api/dashboard")

    assert response.status_code == 200
    dashboard = response.json()
    assert dashboard["active_task_count"] == 2
    assert dashboard["due_soon_count"] == 0
    assert dashboard["overdue_count"] == 1
    assert dashboard["materials_count"] == 1
    assert dashboard["courses_count"] == 1
    assert dashboard["recent_materials"][0]["course_name"] == "操作系统"
    assert dashboard["overdue_tasks"][0]["name"] == "逾期任务"


def test_m5_deleting_material_keeps_task_source_explanation(client):
    material = client.post(
        "/api/materials",
        json={"original_filename": "待删除的通知.txt"},
    ).json()
    task = client.post(
        "/api/tasks",
        json={"material_id": material["id"], "name": "来源任务"},
    ).json()

    deleted = client.delete(f"/api/materials/{material['id']}")

    assert deleted.status_code == 204
    saved = client.get(f"/api/tasks/{task['id']}")
    assert saved.status_code == 200
    assert saved.json()["material_id"] is None
    assert saved.json()["source_material_name"] == "待删除的通知.txt"


def test_required_fields_have_clear_validation_error(client):
    response = client.post("/api/courses", json={})
    assert response.status_code == 422
    body = response.json()
    assert body["error"]["code"] == "VALIDATION_ERROR"
    assert body["error"]["details"]


def test_missing_resource_uses_unified_error_shape(client):
    response = client.get("/api/tasks/999999")
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "TASK_NOT_FOUND"


def test_invalid_related_ids_are_rejected_before_database_write(client):
    material = client.post(
        "/api/materials",
        json={"course_id": 999999, "original_filename": "不存在的课程.txt"},
    )
    assert material.status_code == 404
    assert material.json()["error"]["code"] == "COURSE_NOT_FOUND"

    task = client.post(
        "/api/tasks",
        json={"material_id": 999999, "name": "不存在的资料任务"},
    )
    assert task.status_code == 404
    assert task.json()["error"]["code"] == "MATERIAL_NOT_FOUND"


def test_past_due_unfinished_task_is_marked_overdue(client):
    created = client.post(
        "/api/tasks",
        json={
            "name": "过去的截止任务",
            "due_at": "2020-01-01T12:00:00",
            "status": "not_started",
        },
    )
    assert created.status_code == 201
    task_id = created.json()["id"]

    listed = client.get("/api/tasks")
    assert listed.status_code == 200
    task = next(item for item in listed.json() if item["id"] == task_id)
    assert task["status"] == "overdue"


def test_m3_uploads_and_extracts_utf8_text(client, isolated_upload_dir):
    response = _upload(client, "通知.md", "数据库实验一\n截止时间：10月15日".encode(), "text/markdown")

    assert response.status_code == 201
    material = response.json()[0]
    assert material["processing_status"] == "processed"
    assert "数据库实验一" in material["extracted_text"]
    assert material["file_size"] > 0
    assert len(material["content_hash"]) == 64
    assert (isolated_upload_dir / material["stored_path"]).is_file()


def test_m8_upload_auto_extracts_local_candidates_without_creating_tasks(client, isolated_upload_dir):
    response = _upload(
        client,
        "自动抽取通知.txt",
        "课程：数据库原理\n作业一：完成 SQL 练习，截止时间：2026年10月15日 23:59".encode(),
        "text/plain",
    )

    assert response.status_code == 201
    material = response.json()[0]
    assert material["processing_status"] == "processed"
    assert material["extraction_status"] == "ready"
    assert material["extraction_provider"] == "local-rules"
    assert material["extraction_result"]["tasks"]
    assert material["extraction_error"] is None
    assert client.get("/api/tasks").json() == []
    assert (isolated_upload_dir / material["stored_path"]).is_file()


def test_m8_upload_with_no_candidates_stores_traceable_ready_result(client, isolated_upload_dir):
    response = _upload(
        client,
        "课程范围.txt",
        "课程：数据库原理\n本章重点：索引结构与事务隔离级别".encode(),
        "text/plain",
    )

    assert response.status_code == 201
    material = response.json()[0]
    assert material["extraction_status"] == "ready"
    assert material["extraction_provider"] == "local-rules"
    assert material["extraction_result"]["batch_id"]
    assert material["extraction_result"]["tasks"] == []
    assert material["extraction_result"]["warnings"] == []


def test_m8_local_extraction_failure_does_not_fail_upload(client, isolated_upload_dir, monkeypatch):
    from app.api import materials as materials_api
    from app.services.extraction import ExtractionError

    def fail_local_extraction(material, provider_name=None):
        assert provider_name == "local-rules"
        raise ExtractionError("LOCAL_RULES_FAILED", "本地规则暂时不可用")

    monkeypatch.setattr(materials_api, "extract_material", fail_local_extraction)

    response = _upload(client, "本地抽取失败.txt", "作业截止时间：2026年10月15日 23:59".encode(), "text/plain")

    assert response.status_code == 201
    material = response.json()[0]
    assert material["processing_status"] == "processed"
    assert material["extraction_status"] == "failed"
    assert material["extraction_provider"] == "local-rules"
    assert material["extraction_error"] == "本地规则暂时不可用"
    assert material["extraction_result"] is None


def test_m7_ocr_failure_keeps_original_file(client, isolated_upload_dir, monkeypatch):
    from app.services import file_parser
    from app.services.file_parser import FileProcessingError

    def fail_ocr(_content):
        raise FileProcessingError("OCR_FAILED", "OCR 服务暂时不可用")

    monkeypatch.setattr(file_parser, "_extract_image_ocr", fail_ocr)
    image = BytesIO()
    Image.new("RGB", (80, 40), "white").save(image, format="PNG")

    response = _upload(client, "OCR失败.png", image.getvalue(), "image/png")

    assert response.status_code == 201
    material = response.json()[0]
    assert material["processing_status"] == "failed"
    assert material["processing_error"] == "OCR 服务暂时不可用"
    assert (isolated_upload_dir / material["stored_path"]).is_file()


def test_m3_uploads_and_extracts_docx(client, isolated_upload_dir):
    document = Document()
    document.add_paragraph("数据结构作业要求")
    document.add_paragraph("请完成二叉树遍历练习")
    content = BytesIO()
    document.save(content)

    response = _upload(
        client,
        "作业要求.docx",
        content.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert response.status_code == 201
    assert response.json()[0]["processing_status"] == "processed"
    assert "二叉树遍历" in response.json()[0]["extracted_text"]


def test_m3_uploads_and_extracts_pdf(client, isolated_upload_dir):
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Calculus exam notice")
    content = document.tobytes()
    document.close()

    response = _upload(client, "考试通知.pdf", content, "application/pdf")

    assert response.status_code == 201
    assert response.json()[0]["processing_status"] == "processed"
    assert "Calculus exam notice" in response.json()[0]["extracted_text"]


def test_m3_scanned_pdf_uses_ocr_fallback(client, isolated_upload_dir):
    image = Image.new("RGB", (420, 80), "white")
    ImageDraw.Draw(image).text((10, 20), "Assignment due Friday", fill="black")
    image_bytes = BytesIO()
    image.save(image_bytes, format="PNG")

    document = fitz.open()
    page = document.new_page(width=420, height=80)
    page.insert_image(fitz.Rect(0, 0, 420, 80), stream=image_bytes.getvalue())
    content = document.tobytes()
    document.close()

    response = _upload(client, "扫描通知.pdf", content, "application/pdf")

    assert response.status_code == 201
    material = response.json()[0]
    assert material["processing_status"] == "processed"
    assert "AssignmentdueFriday" in material["extracted_text"].replace(" ", "")


def test_m3_exposes_upload_policy(client):
    response = client.get("/api/materials/upload-policy")

    assert response.status_code == 200
    policy = response.json()
    assert policy["max_upload_size_mb"] == 20
    assert policy["max_upload_files"] == 10
    assert "pdf" in policy["extensions"]


def test_m3_rejects_unsupported_and_empty_files_without_material_records(client, isolated_upload_dir):
    unsupported = _upload(client, "程序.exe", b"not allowed", "application/octet-stream")
    empty = _upload(client, "空文件.txt", b"", "text/plain")

    assert unsupported.status_code == 400
    assert unsupported.json()["error"]["code"] == "UNSUPPORTED_FILE_TYPE"
    assert empty.status_code == 400
    assert empty.json()["error"]["code"] == "EMPTY_FILE"
    assert client.get("/api/materials").json() == []
    assert list(isolated_upload_dir.iterdir()) == []


def test_m3_detects_duplicate_content(client, isolated_upload_dir):
    content = b"same content"
    first = _upload(client, "第一份.txt", content, "text/plain")
    second = _upload(client, "第二份.txt", content, "text/plain")

    assert first.status_code == 201
    assert second.status_code == 409
    assert second.json()["error"]["code"] == "DUPLICATE_FILE"
    assert len(client.get("/api/materials").json()) == 1


def test_m3_keeps_failed_original_file_and_allows_retry(client, isolated_upload_dir):
    failed = _upload(client, "扫描通知.pdf", b"%PDF-not-a-real-pdf", "application/pdf")
    material = failed.json()[0]
    stored_path = isolated_upload_dir / material["stored_path"]

    assert failed.status_code == 201
    assert material["processing_status"] == "failed"
    assert material["processing_error"]
    assert stored_path.is_file()

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Retry extraction works")
    stored_path.write_bytes(document.tobytes())
    document.close()

    retried = client.post(f"/api/materials/{material['id']}/retry")
    assert retried.status_code == 200
    retried_material = retried.json()
    assert retried_material["processing_status"] == "processed"
    assert "Retry extraction works" in retried_material["extracted_text"]
    assert retried_material["extraction_status"] == "ready"
    assert retried_material["extraction_provider"] == "local-rules"
    assert retried_material["extraction_result"]["tasks"] == []


def test_m3_manual_create_cannot_set_server_managed_file_fields(client):
    response = client.post(
        "/api/materials",
        json={
            "original_filename": "手动资料.txt",
            "stored_path": "../../outside.txt",
            "content_hash": "a" * 64,
            "processing_status": "processed",
        },
    )

    assert response.status_code == 201
    material = response.json()
    assert material["stored_path"] is None
    assert material["content_hash"] is None
    assert material["processing_status"] == "pending"


def test_m3_delete_removes_original_file(client, isolated_upload_dir):
    uploaded = _upload(client, "待删除.txt", b"will be deleted", "text/plain")
    material = uploaded.json()[0]
    stored_path = isolated_upload_dir / material["stored_path"]
    assert stored_path.is_file()

    deleted = client.delete(f"/api/materials/{material['id']}")

    assert deleted.status_code == 204
    assert not stored_path.exists()


def test_m3_retry_uses_stored_type_after_original_name_changes(client, isolated_upload_dir):
    failed = _upload(client, "原始通知.pdf", b"%PDF-invalid", "application/pdf")
    material = failed.json()[0]
    stored_path = isolated_upload_dir / material["stored_path"]

    renamed = client.patch(
        f"/api/materials/{material['id']}",
        json={"original_filename": "用户改名.txt"},
    )
    assert renamed.status_code == 200

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Retry still detects PDF")
    stored_path.write_bytes(document.tobytes())
    document.close()

    retried = client.post(f"/api/materials/{material['id']}/retry")

    assert retried.status_code == 200
    assert retried.json()["processing_status"] == "processed"
    assert "Retry still detects PDF" in retried.json()["extracted_text"]


def test_m3_manual_text_marks_failed_material_as_processed(client, isolated_upload_dir):
    failed = _upload(client, "需要补录.txt", b"\xff\xfe\x00", "text/plain")
    material = failed.json()[0]
    assert material["processing_status"] == "failed"

    updated = client.patch(
        f"/api/materials/{material['id']}",
        json={"extracted_text": "用户补充的正文"},
    )

    assert updated.status_code == 200
    assert updated.json()["processing_status"] == "processed"
    assert updated.json()["processing_error"] is None


def test_m3_uploaded_file_type_cannot_be_changed(client, isolated_upload_dir):
    uploaded = _upload(client, "类型固定.txt", b"plain text", "text/plain")
    material = uploaded.json()[0]

    updated = client.patch(
        f"/api/materials/{material['id']}",
        json={"file_type": "pdf"},
    )

    assert updated.status_code == 409
    assert updated.json()["error"]["code"] == "FILE_TYPE_IMMUTABLE"


def test_m3_rejects_too_many_files_before_reading(client, isolated_upload_dir):
    files = [
        ("files", (f"file-{index}.txt", b"content", "text/plain"))
        for index in range(11)
    ]

    response = client.post("/api/materials/upload", files=files)

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "TOO_MANY_FILES"
    assert client.get("/api/materials").json() == []
    assert list(isolated_upload_dir.iterdir()) == []


def test_m7_upload_strips_path_components_and_enforces_size_limit(client, isolated_upload_dir):
    safe_name = _upload(client, "..\\outside\\安全资料.txt", b"safe text", "text/plain")
    assert safe_name.status_code == 201
    saved = safe_name.json()[0]
    assert saved["original_filename"] == "安全资料.txt"
    assert (isolated_upload_dir / saved["stored_path"]).is_file()

    from app.config import settings

    previous_limit = settings.max_upload_size_mb
    object.__setattr__(settings, "max_upload_size_mb", 0)
    try:
        oversized = _upload(client, "too-large.txt", b"x", "text/plain")
    finally:
        object.__setattr__(settings, "max_upload_size_mb", previous_limit)

    assert oversized.status_code == 400
    assert oversized.json()["error"]["code"] == "FILE_TOO_LARGE"


def test_m4_extracts_structured_ddl_from_material_text(client):
    material = client.post(
        "/api/materials",
        json={
            "original_filename": "数据结构通知.md",
            "extracted_text": "课程：数据结构\n作业一：完成二叉树遍历，截止时间：2026年10月15日 23:59",
        },
    ).json()

    response = client.post(f"/api/materials/{material['id']}/extract")

    assert response.status_code == 200
    result = response.json()
    assert result["status"] == "ready"
    assert result["provider"] == "local-rules"
    assert result["course_name"] == "数据结构"
    assert result["tasks"][0]["due_at"] == "2026-10-15T15:59:00Z"
    assert result["tasks"][0]["source_quote"] in material["extracted_text"]


def test_m8_local_rules_combines_filename_and_body_without_fabricating_quote(client):
    body = "2026年10月15日 23:59"
    material = client.post(
        "/api/materials",
        json={"original_filename": "数据结构-作业通知.txt", "extracted_text": body},
    ).json()

    response = client.post(f"/api/materials/{material['id']}/extract", json={"provider": "local-rules"})

    assert response.status_code == 200
    result = response.json()
    task = result["tasks"][0]
    assert result["course_name"] == "数据结构"
    assert task["name"] == "作业通知"
    assert task["task_type"] == "作业"
    assert task["due_at"] == "2026-10-15T15:59:00Z"
    assert task["source_quote"] == body
    assert task["source_quote"] in material["extracted_text"]
    assert "作业" in result["tags"]

    ddl_material = client.post(
        "/api/materials",
        json={"original_filename": "数据结构-DDL通知.txt", "extracted_text": body},
    ).json()
    ddl_result = client.post(f"/api/materials/{ddl_material['id']}/extract").json()
    assert ddl_result["tasks"][0]["task_type"] == "DDL"
    assert "DDL" in ddl_result["tags"]

    ambiguous = client.post(
        "/api/materials",
        json={"original_filename": "作业通知.txt", "extracted_text": body},
    ).json()
    ambiguous_result = client.post(f"/api/materials/{ambiguous['id']}/extract").json()
    assert ambiguous_result["course_name"] is None


def test_m8_extraction_policy_reports_local_and_external_availability(client):
    from app.config import settings

    previous_key, previous_model = settings.llm_api_key, settings.llm_model
    object.__setattr__(settings, "llm_api_key", "")
    object.__setattr__(settings, "llm_model", "")
    try:
        response = client.get("/api/materials/extraction-policy")
    finally:
        object.__setattr__(settings, "llm_api_key", previous_key)
        object.__setattr__(settings, "llm_model", previous_model)

    assert response.status_code == 200
    result = response.json()
    assert result["default_provider"] == "local-rules"
    providers = {provider["id"]: provider for provider in result["providers"]}
    assert providers["local-rules"]["available"] is True
    assert providers["local-rules"]["sends_data_externally"] is False
    assert providers["openai-compatible"]["available"] is False
    assert providers["openai-compatible"]["sends_data_externally"] is True


def test_m8_external_extraction_requires_server_configuration(client):
    from app.config import settings

    previous_key, previous_model = settings.llm_api_key, settings.llm_model
    object.__setattr__(settings, "llm_api_key", "")
    object.__setattr__(settings, "llm_model", "")
    try:
        material = client.post(
            "/api/materials",
            json={"original_filename": "外部模型通知.txt", "extracted_text": "作业截止时间：2026年10月15日"},
        ).json()
        response = client.post(
            f"/api/materials/{material['id']}/extract",
            json={"provider": "openai-compatible"},
        )
    finally:
        object.__setattr__(settings, "llm_api_key", previous_key)
        object.__setattr__(settings, "llm_model", previous_model)

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "LLM_NOT_CONFIGURED"


def test_m8_can_explicitly_select_external_extraction(client, monkeypatch):
    from app.config import settings
    from app.services.llm_provider import OpenAICompatibleProvider

    previous_key, previous_model = settings.llm_api_key, settings.llm_model
    object.__setattr__(settings, "llm_api_key", "test-key")
    object.__setattr__(settings, "llm_model", "test-model")
    monkeypatch.setattr(
        OpenAICompatibleProvider,
        "extract",
        lambda self, text, filename: {
            "course_name": "机器学习",
            "material_type": "作业要求",
            "tags": ["作业"],
            "tasks": [
                {
                    "name": "完成模型实验",
                    "task_type": "实验",
                    "due_at": "2026-10-15 18:00",
                    "source_quote": text,
                    "confidence": 0.92,
                }
            ],
        },
    )
    try:
        material = client.post(
            "/api/materials",
            json={"original_filename": "机器学习通知.txt", "extracted_text": "完成模型实验，截止时间：2026年10月15日 18:00"},
        ).json()
        response = client.post(
            f"/api/materials/{material['id']}/extract",
            json={"provider": "openai-compatible"},
        )
    finally:
        object.__setattr__(settings, "llm_api_key", previous_key)
        object.__setattr__(settings, "llm_model", previous_model)

    assert response.status_code == 200
    result = response.json()
    assert result["provider"] == "openai-compatible"
    assert result["course_name"] == "机器学习"
    assert result["tasks"][0]["name"] == "完成模型实验"


def test_m4_missing_date_requires_human_review(client):
    material = client.post(
        "/api/materials",
        json={"original_filename": "实验要求.txt", "extracted_text": "请完成实验一并提交实验报告"},
    ).json()

    result = client.post(f"/api/materials/{material['id']}/extract").json()

    assert result["status"] == "needs_review"
    assert result["needs_review"] is True
    assert "MISSING_DUE_DATE" in result["warnings"]
    assert result["tasks"][0]["need_review"] is True


def test_m4_detects_multiple_dates_in_one_source_quote(client):
    material = client.post(
        "/api/materials",
        json={
            "original_filename": "考试安排.txt",
            "extracted_text": "课程：算法\n考试安排：2026年10月15日或2026年10月20日截止，请确认最终日期",
        },
    ).json()

    result = client.post(f"/api/materials/{material['id']}/extract").json()

    assert result["status"] == "needs_review"
    assert "MULTIPLE_DATES_IN_SOURCE" in result["warnings"]
    assert "MULTIPLE_DATES_IN_SOURCE" in result["tasks"][0]["warnings"]


def test_m4_invalid_provider_json_is_saved_as_failed_extraction(client, monkeypatch):
    from app.services import extraction

    monkeypatch.setattr(extraction, "get_provider", lambda: _BrokenProvider())
    material = client.post(
        "/api/materials",
        json={"original_filename": "模型结果.txt", "extracted_text": "作业一截止2026年10月15日"},
    ).json()

    response = client.post(f"/api/materials/{material['id']}/extract")

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "INVALID_PROVIDER_JSON"
    saved = client.get(f"/api/materials/{material['id']}").json()
    assert saved["extraction_status"] == "failed"
    assert saved["extraction_error"]


def test_m7_provider_timeout_is_saved_as_a_safe_failure(client, monkeypatch):
    from app.services import extraction

    monkeypatch.setattr(extraction, "get_provider", lambda: _TimeoutProvider())
    material = client.post(
        "/api/materials",
        json={"original_filename": "超时通知.txt", "extracted_text": "作业截止时间：2026年10月15日"},
    ).json()

    response = client.post(f"/api/materials/{material['id']}/extract")

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "LLM_TIMEOUT"
    saved = client.get(f"/api/materials/{material['id']}").json()
    assert saved["extraction_status"] == "failed"
    assert saved["extraction_error"] == "模型调用超时，请稍后重试"


def test_m7_provider_network_failure_is_saved_as_a_safe_failure(client, monkeypatch):
    from app.services import extraction

    monkeypatch.setattr(extraction, "get_provider", lambda: _NetworkFailureProvider())
    material = client.post(
        "/api/materials",
        json={"original_filename": "网络失败通知.txt", "extracted_text": "作业截止时间：2026年10月15日"},
    ).json()

    response = client.post(f"/api/materials/{material['id']}/extract")

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "LLM_REQUEST_FAILED"
    saved = client.get(f"/api/materials/{material['id']}").json()
    assert saved["extraction_status"] == "failed"


def test_m4_confirmation_creates_tasks_only_after_user_confirmation(client):
    course = client.post("/api/courses", json={"name": "数据库原理"}).json()
    material = client.post(
        "/api/materials",
        json={
            "course_id": course["id"],
            "original_filename": "DDL通知.txt",
            "extracted_text": "课程：数据库原理\n作业一：完成SQL练习，截止时间：2026年10月15日",
        },
    ).json()

    extracted = client.post(f"/api/materials/{material['id']}/extract").json()
    assert client.get("/api/tasks").json() == []
    candidate = extracted["tasks"][0]
    candidate["name"] = "人工确认后的 SQL 练习"

    confirmed = client.post(
        f"/api/materials/{material['id']}/extraction/confirm",
        json={"tasks": [candidate]},
    )

    assert confirmed.status_code == 200
    assert confirmed.json()["status"] == "confirmed"
    assert confirmed.json()["confirmed_task_ids"]
    tasks = client.get("/api/tasks").json()
    assert len(tasks) == 1
    assert tasks[0]["name"] == "人工确认后的 SQL 练习"
    assert tasks[0]["course_id"] == course["id"]
    assert tasks[0]["need_review"] is False


def test_m4_cannot_reextract_confirmed_material_or_destroy_its_result(client):
    material = client.post(
        "/api/materials",
        json={
            "original_filename": "已确认通知.txt",
            "extracted_text": "作业一：完成 SQL 练习，截止时间：2026年10月15日 23:59",
        },
    ).json()
    extracted = client.post(f"/api/materials/{material['id']}/extract").json()
    confirmed = client.post(
        f"/api/materials/{material['id']}/extraction/confirm",
        json={"tasks": extracted["tasks"]},
    )
    assert confirmed.status_code == 200

    response = client.post(f"/api/materials/{material['id']}/extract")

    assert response.status_code == 409
    assert response.json()["error"]["code"] == "EXTRACTION_ALREADY_CONFIRMED"
    saved = client.get(f"/api/materials/{material['id']}").json()
    assert saved["extraction_status"] == "confirmed"
    assert saved["extraction_result"]


def test_m8_retry_cannot_clear_confirmed_extraction_snapshot(client):
    material = client.post(
        "/api/materials",
        json={
            "original_filename": "已确认且无文件.txt",
            "extracted_text": "作业一：完成 SQL 练习，截止时间：2026年10月15日 23:59",
        },
    ).json()
    extracted = client.post(f"/api/materials/{material['id']}/extract").json()
    confirmed = client.post(
        f"/api/materials/{material['id']}/extraction/confirm",
        json={"tasks": extracted["tasks"]},
    )
    assert confirmed.status_code == 200
    before = client.get(f"/api/materials/{material['id']}").json()

    response = client.post(f"/api/materials/{material['id']}/retry")

    assert response.status_code == 409
    assert response.json()["error"]["code"] == "EXTRACTION_ALREADY_CONFIRMED"
    after = client.get(f"/api/materials/{material['id']}").json()
    assert after["extraction_status"] == "confirmed"
    assert after["extraction_result"] == before["extraction_result"]


def test_m4_editing_metadata_does_not_reset_extraction_state(client):
    material = client.post(
        "/api/materials",
        json={
            "original_filename": "可编辑通知.txt",
            "extracted_text": "作业一：完成 SQL 练习，截止时间：2026年10月15日 23:59",
        },
    ).json()
    extracted = client.post(f"/api/materials/{material['id']}/extract").json()
    assert extracted["status"] == "ready"

    updated = client.patch(
        f"/api/materials/{material['id']}",
        json={"original_filename": "改名后的通知.txt", "extracted_text": material["extracted_text"]},
    )

    assert updated.status_code == 200
    assert updated.json()["extraction_status"] == "ready"
    assert updated.json()["extraction_result"]


def test_m4_confirmation_persists_edited_snapshot_and_material_metadata(client):
    course = client.post("/api/courses", json={"name": "程序设计"}).json()
    material = client.post(
        "/api/materials",
        json={
            "original_filename": "程序设计通知.txt",
            "extracted_text": "作业一：完成程序设计报告，截止时间：2026年10月15日 23:59",
        },
    ).json()
    extracted = client.post(f"/api/materials/{material['id']}/extract").json()
    candidate = extracted["tasks"][0]
    candidate["name"] = "人工修订后的程序设计报告"

    confirmed = client.post(
        f"/api/materials/{material['id']}/extraction/confirm",
        json={
            "course_id": course["id"],
            "material_type": "作业要求",
            "tags": ["作业", "重点"],
            "tasks": [candidate],
        },
    )

    assert confirmed.status_code == 200
    assert confirmed.json()["material_type"] == "作业要求"
    assert confirmed.json()["tags"] == ["作业", "重点"]
    assert confirmed.json()["tasks"][0]["name"] == "人工修订后的程序设计报告"
    saved_material = client.get(f"/api/materials/{material['id']}").json()
    assert saved_material["material_type"] == "作业要求"
    assert saved_material["tags"] == ["作业", "重点"]
    saved_tasks = client.get("/api/tasks").json()
    assert saved_tasks[0]["course_id"] == course["id"]


def test_m4_parses_english_month_dates(client):
    material = client.post(
        "/api/materials",
        json={
            "original_filename": "english-notice.txt",
            "extracted_text": "Assignment due October 15, 2026",
        },
    ).json()

    result = client.post(f"/api/materials/{material['id']}/extract").json()

    assert result["tasks"][0]["due_at"] == "2026-10-15T15:59:00Z"
    assert "MISSING_DUE_DATE" not in result["tasks"][0]["warnings"]


def test_m4_parses_relative_weekday_and_chinese_pm_time(client):
    material = client.post(
        "/api/materials",
        json={
            "original_filename": "相对日期通知.txt",
            "extracted_text": "实验一：请在下周五下午5点前提交实验报告",
        },
    ).json()

    result = client.post(f"/api/materials/{material['id']}/extract").json()

    task = result["tasks"][0]
    assert task["due_at"].endswith("T09:00:00Z")
    assert "MISSING_DUE_DATE" not in task["warnings"]
    assert "INVALID_TIME" not in task["warnings"]


def test_m4_confirmed_candidate_cannot_be_confirmed_twice(client):
    material = client.post(
        "/api/materials",
        json={
            "original_filename": "幂等确认通知.txt",
            "extracted_text": "作业一：完成 SQL 练习，截止时间：2026年10月15日 23:59",
        },
    ).json()
    extracted = client.post(f"/api/materials/{material['id']}/extract").json()
    first = client.post(
        f"/api/materials/{material['id']}/extraction/confirm",
        json={"tasks": extracted["tasks"]},
    )
    second = client.post(
        f"/api/materials/{material['id']}/extraction/confirm",
        json={"tasks": extracted["tasks"]},
    )

    assert first.status_code == 200
    assert second.status_code == 409
    assert second.json()["error"]["code"] == "EXTRACTION_ALREADY_CONFIRMED"
    assert len(client.get("/api/tasks").json()) == 1


def test_m6_generates_editable_staged_plan_and_preserves_manual_edits(client):
    course = client.post("/api/courses", json={"name": "专利评估与投资"}).json()
    material = client.post(
        "/api/materials",
        json={
            "course_id": course["id"],
            "original_filename": "实验重点.md",
            "tags": ["专利评估", "投资决策"],
            "summary": "掌握专利价值分析与投资决策方法",
        },
    ).json()
    task = client.post(
        "/api/tasks",
        json={
            "course_id": course["id"],
            "material_id": material["id"],
            "name": "完成实验报告",
            "description": "整理实验结论并检查格式",
        },
    ).json()
    exam_date = date.today() + timedelta(days=5)

    response = client.post(
        "/api/study-plans/generate",
        json={"course_id": course["id"], "exam_date": exam_date.isoformat(), "daily_minutes": 60},
    )

    assert response.status_code == 201
    plan = response.json()
    assert len(plan["items"]) == 2
    assert plan["material_count"] == 1
    assert plan["task_count"] == 1
    assert plan["completed_item_count"] == 0
    assert plan["total_minutes"] == 90
    assert plan["completed_minutes"] == 0
    assert any("专利评估" in point for point in plan["items"][0]["knowledge_points"])
    assert any(item["source_task_ids"] == [task["id"]] for item in plan["items"])
    assert all(item["date"] <= exam_date.isoformat() for item in plan["items"])

    edited_items = plan["items"]
    edited_items[0]["content"] = "用户已经重新安排：完成专利价值分析笔记。"
    edited_items[0]["status"] = "completed"
    updated = client.patch(f"/api/study-plans/{plan['id']}", json={"items": edited_items})
    assert updated.status_code == 200
    assert updated.json()["items"][0]["content"].startswith("用户已经重新安排")
    assert updated.json()["items"][0]["status"] == "completed"
    assert updated.json()["completed_item_count"] == 1
    assert updated.json()["completed_minutes"] == 60

    regenerated = client.post(
        "/api/study-plans/generate",
        json={"course_id": course["id"], "exam_date": exam_date.isoformat(), "daily_minutes": 90},
    )
    assert regenerated.status_code == 201
    saved = client.get(f"/api/study-plans/{plan['id']}")
    assert saved.json()["items"][0]["content"].startswith("用户已经重新安排")

    archived = client.post(f"/api/study-plans/{plan['id']}/archive")
    assert archived.status_code == 200
    assert archived.json()["status"] == "archived"
    assert all(item["id"] != plan["id"] for item in client.get("/api/study-plans").json())
    assert client.get("/api/study-plans", params={"include_archived": True}).json()


def test_m6_rejects_plan_item_after_exam_date(client):
    course = client.post("/api/courses", json={"name": "算法设计"}).json()
    material = client.post(
        "/api/materials",
        json={"course_id": course["id"], "original_filename": "算法复习范围.md"},
    )
    assert material.status_code == 201
    exam_date = date.today() + timedelta(days=2)
    plan = client.post(
        "/api/study-plans/generate",
        json={"course_id": course["id"], "exam_date": exam_date.isoformat(), "daily_minutes": 45},
    ).json()
    invalid_items = plan["items"]
    invalid_items[0]["date"] = (exam_date + timedelta(days=1)).isoformat()

    response = client.patch(f"/api/study-plans/{plan['id']}", json={"items": invalid_items})

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "PLAN_DATE_AFTER_EXAM"


def test_m6_reports_topics_that_do_not_fit_available_minutes(client):
    course = client.post("/api/courses", json={"name": "操作系统"}).json()
    for index in range(4):
        client.post(
            "/api/materials",
            json={
                "course_id": course["id"],
                "original_filename": f"章节重点-{index}.txt",
                "tags": [f"知识点-{index}"],
            },
        )

    response = client.post(
        "/api/study-plans/generate",
        json={
            "course_id": course["id"],
            "exam_date": (date.today() + timedelta(days=1)).isoformat(),
            "daily_minutes": 60,
        },
    )

    assert response.status_code == 201
    plan = response.json()
    assert any("多个主题" in warning for warning in plan["warnings"])
    assert sum(len(item["source_material_ids"]) for item in plan["items"]) == 2
    assert sum("资料“" in warning and "60 分钟未排入" in warning for warning in plan["warnings"]) == 2
    assert all(item["minutes"] <= 60 for item in plan["items"])


def test_m7_empty_course_plan_explains_missing_sources(client):
    course = client.post("/api/courses", json={"name": "空课程"}).json()

    response = client.post(
        "/api/study-plans/generate",
        json={
            "course_id": course["id"],
            "exam_date": (date.today() + timedelta(days=2)).isoformat(),
            "daily_minutes": 30,
        },
    )

    assert response.status_code == 201
    assert any("没有资料或未完成任务" in warning for warning in response.json()["warnings"])


def test_m6_exports_tasks_materials_and_study_plan(client):
    course = client.post("/api/courses", json={"name": "数据库原理"}).json()
    material = client.post(
        "/api/materials",
        json={
            "course_id": course["id"],
            "original_filename": "数据库重点.txt",
            "material_type": "复习资料",
            "tags": ["索引", "事务"],
            "summary": "索引结构和事务隔离级别",
        },
    ).json()
    client.post(
        "/api/tasks",
        json={
            "course_id": course["id"],
            "material_id": material["id"],
            "name": "数据库练习题",
            "due_at": "2099-01-02T12:00:00",
            "description": "完成练习并提交",
        },
    )
    plan = client.post(
        "/api/study-plans/generate",
        json={
            "course_id": course["id"],
            "exam_date": (date.today() + timedelta(days=1)).isoformat(),
            "daily_minutes": 30,
        },
    ).json()

    tasks_csv = client.get("/api/exports/tasks.csv")
    materials_md = client.get("/api/exports/materials.md")
    plan_md = client.get(f"/api/exports/study-plans/{plan['id']}.md")

    assert tasks_csv.status_code == 200
    assert tasks_csv.content.startswith(b"\xef\xbb\xbf")
    assert "数据库练习题" in tasks_csv.content.decode("utf-8")
    assert 'filename="ddl-tasks.csv"' in tasks_csv.headers["content-disposition"]
    assert materials_md.status_code == 200
    assert "数据库重点.txt" in materials_md.text
    assert "事务" in materials_md.text
    assert plan_md.status_code == 200
    assert "数据库原理复习计划" in plan_md.text
    assert "复习清单" in plan_md.text
    assert "数据库重点.txt" in plan_md.text
    assert "数据库练习题" in plan_md.text


def test_tasks_icalendar_export_has_timezone_alarm_and_active_filter(client):
    course = client.post("/api/courses", json={"name": "操作系统"}).json()
    active = client.post(
        "/api/tasks",
        json={
            "course_id": course["id"],
            "name": "提交进程调度实验",
            "task_type": "实验",
            "description": "检查报告和源代码",
            "due_at": "2026-09-01T20:30:00",
            "priority": 5,
        },
    ).json()
    client.post(
        "/api/tasks",
        json={"course_id": course["id"], "name": "已完成作业", "due_at": "2026-09-02T20:30:00", "status": "completed"},
    )
    client.post("/api/tasks", json={"course_id": course["id"], "name": "待确定日期的任务"})

    response = client.get("/api/exports/tasks.ics")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/calendar")
    assert 'filename="ddl-tasks.ics"' in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    calendar = Calendar.from_ical(response.content)
    assert len(calendar.walk("VTIMEZONE")) == 1
    events = calendar.walk("VEVENT")
    assert len(events) == 1
    event = events[0]
    assert str(event["uid"]) == f"task-{active['navigation_key']}@learning-assistant.local"
    assert str(event["summary"]) == "[DDL] 提交进程调度实验"
    assert event.decoded("dtstart").utcoffset() == timedelta(hours=8)
    assert "检查报告和源代码" in str(event["description"])
    alarms = event.walk("VALARM")
    assert len(alarms) == 1
    assert alarms[0].decoded("trigger") == timedelta(days=-1)

    with_completed = client.get("/api/exports/tasks.ics", params={"include_completed": True})
    assert len(Calendar.from_ical(with_completed.content).walk("VEVENT")) == 2


def test_m7_demo_seed_is_idempotent():
    from sqlalchemy import create_engine, select
    from sqlalchemy.orm import Session
    from sqlalchemy.pool import StaticPool

    from app.database import Base
    from app.models.course import Course
    from app.models.material import Material
    from app.models.task import Task
    from app.services.demo_data import ensure_demo_data

    demo_engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(bind=demo_engine)
    try:
        with Session(demo_engine) as db:
            assert ensure_demo_data(db) is True
            assert ensure_demo_data(db) is False
            assert db.scalar(select(Course).where(Course.name == "数据结构（演示）")) is not None
            assert db.query(Material).count() == 2
            assert db.query(Task).count() == 2
    finally:
        demo_engine.dispose()
