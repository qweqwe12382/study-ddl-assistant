"""Validation and text extraction for uploaded learning materials."""

from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from threading import Lock
import zipfile

import fitz
from docx import Document

from app.config import settings


ALLOWED_FILE_TYPES: dict[str, set[str]] = {
    "pdf": {"application/pdf"},
    "docx": {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/zip",
    },
    "txt": {"text/plain"},
    "md": {"text/markdown", "text/plain"},
    "png": {"image/png"},
    "jpg": {"image/jpeg"},
    "jpeg": {"image/jpeg"},
    "gif": {"image/gif"},
    "bmp": {"image/bmp"},
    "webp": {"image/webp"},
}
GENERIC_MIME_TYPES = {"", "application/octet-stream"}
IMAGE_TYPES = {"png", "jpg", "jpeg", "gif", "bmp", "webp"}
MAX_IMAGE_PIXELS = 25_000_000
MAX_PDF_OCR_PAGES = 20
_OCR_LOCK = Lock()


class FileProcessingError(ValueError):
    """An expected upload or parsing failure with a user-facing message."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class PreparedUpload:
    original_filename: str
    file_type: str
    content_type: str
    content: bytes
    content_hash: str


def _clean_filename(filename: str | None) -> str:
    if not filename:
        raise FileProcessingError("INVALID_FILENAME", "文件名不能为空")
    cleaned = Path(filename.replace("\\", "/")).name.replace("\x00", "").strip()
    if not cleaned or cleaned in {".", ".."}:
        raise FileProcessingError("INVALID_FILENAME", "文件名无效")
    return cleaned[:255]


def prepare_upload(filename: str | None, content_type: str | None, content: bytes) -> PreparedUpload:
    original_filename = _clean_filename(filename)
    suffix = Path(original_filename).suffix.lower().lstrip(".")
    if suffix not in ALLOWED_FILE_TYPES:
        raise FileProcessingError("UNSUPPORTED_FILE_TYPE", "仅支持 PDF、DOCX、TXT、MD 和常见图片格式")
    if not content:
        raise FileProcessingError("EMPTY_FILE", "不能上传空文件")
    max_bytes = settings.max_upload_size_mb * 1024 * 1024
    if len(content) > max_bytes:
        raise FileProcessingError("FILE_TOO_LARGE", f"文件大小不能超过 {settings.max_upload_size_mb} MB")

    normalized_mime = (content_type or "").lower().split(";", 1)[0].strip()
    if normalized_mime not in GENERIC_MIME_TYPES and normalized_mime not in ALLOWED_FILE_TYPES[suffix]:
        raise FileProcessingError("MIME_TYPE_MISMATCH", "文件扩展名与 MIME 类型不匹配")
    _validate_signature(suffix, content)

    import hashlib

    return PreparedUpload(
        original_filename=original_filename,
        file_type=suffix,
        content_type=normalized_mime,
        content=content,
        content_hash=hashlib.sha256(content).hexdigest(),
    )


def _validate_signature(file_type: str, content: bytes) -> None:
    if file_type == "pdf" and not content.startswith(b"%PDF-"):
        raise FileProcessingError("CORRUPTED_FILE", "文件不是有效的 PDF")
    if file_type == "docx":
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                if "word/document.xml" not in archive.namelist():
                    raise FileProcessingError("CORRUPTED_FILE", "文件不是有效的 DOCX")
                if len(archive.infolist()) > 2_000:
                    raise FileProcessingError("FILE_TOO_LARGE", "DOCX 内部文件数量过多")
                if sum(item.file_size for item in archive.infolist()) > settings.max_upload_size_mb * 1024 * 1024 * 5:
                    raise FileProcessingError("FILE_TOO_LARGE", "DOCX 解压后的内容过大")
        except zipfile.BadZipFile as exc:
            raise FileProcessingError("CORRUPTED_FILE", "文件不是有效的 DOCX") from exc
    if file_type in IMAGE_TYPES:
        try:
            from PIL import Image

            with Image.open(BytesIO(content)) as image:
                if image.width * image.height > MAX_IMAGE_PIXELS:
                    raise FileProcessingError("IMAGE_TOO_LARGE", "图片分辨率过大，无法安全处理")
                image.verify()
        except FileProcessingError:
            raise
        except Exception as exc:
            raise FileProcessingError("CORRUPTED_FILE", "图片文件损坏或格式无法识别") from exc


def extract_text(file_type: str, content: bytes) -> str:
    try:
        if file_type in {"txt", "md"}:
            return _decode_text(content)
        if file_type == "pdf":
            return _extract_pdf(content)
        if file_type == "docx":
            return _extract_docx(content)
        if file_type in IMAGE_TYPES:
            return _extract_image_ocr(content)
    except FileProcessingError:
        raise
    except Exception as exc:
        raise FileProcessingError("PARSING_FAILED", "文件解析失败，请稍后重试或手动录入") from exc
    raise FileProcessingError("UNSUPPORTED_FILE_TYPE", "暂不支持该文件格式")


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        if "\x00" in text and encoding != "utf-16":
            continue
        if text.count("\ufffd") > max(1, len(text) // 100):
            continue
        visible_chars = sum(char.isprintable() or char in "\r\n\t" for char in text)
        if text and visible_chars / len(text) < 0.85:
            continue
        return text.strip()
    raise FileProcessingError("INVALID_TEXT_ENCODING", "文本编码无法识别，请另存为 UTF-8 后重试")


def _extract_pdf(content: bytes) -> str:
    try:
        document = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise FileProcessingError("CORRUPTED_FILE", "PDF 文件损坏或无法打开") from exc
    try:
        text = "\n".join(page.get_text("text") for page in document).strip()
        if text:
            return text

        if len(document) > MAX_PDF_OCR_PAGES:
            raise FileProcessingError(
                "PDF_TOO_MANY_OCR_PAGES",
                f"扫描版 PDF 页数过多，暂只支持不超过 {MAX_PDF_OCR_PAGES} 页的 OCR",
            )

        ocr_chunks: list[str] = []
        for page in document:
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            try:
                ocr_text = _extract_image_ocr(pixmap.tobytes("png"))
            except FileProcessingError as error:
                if error.code == "OCR_NO_TEXT":
                    continue
                raise
            if ocr_text:
                ocr_chunks.append(ocr_text)
        text = "\n".join(ocr_chunks).strip()
    finally:
        document.close()
    if not text:
        raise FileProcessingError("NO_TEXT_FOUND", "PDF 没有可提取文字，可能是扫描版文件")
    return text


def _extract_docx(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise FileProcessingError("CORRUPTED_FILE", "DOCX 文件损坏或无法打开") from exc
    chunks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                chunks.append("\t".join(values))
    text = "\n".join(chunks).strip()
    if not text:
        raise FileProcessingError("NO_TEXT_FOUND", "DOCX 中没有可提取文字")
    return text


@lru_cache(maxsize=1)
def _ocr_engine():
    try:
        from rapidocr_onnxruntime import RapidOCR

        return RapidOCR()
    except Exception as exc:
        raise FileProcessingError("OCR_UNAVAILABLE", "OCR 服务未能启动，请改用文本文件或手动录入") from exc


def _extract_image_ocr(content: bytes) -> str:
    try:
        with _OCR_LOCK:
            result, _ = _ocr_engine()(content)
    except FileProcessingError:
        raise
    except Exception as exc:
        raise FileProcessingError("OCR_FAILED", "图片 OCR 失败，请重试或手动录入") from exc
    lines: list[str] = []
    for item in result or []:
        if len(item) >= 2 and str(item[1]).strip():
            lines.append(str(item[1]).strip())
    text = "\n".join(lines).strip()
    if not text:
        raise FileProcessingError("OCR_NO_TEXT", "图片中没有识别到文字，请上传更清晰的图片")
    return text
